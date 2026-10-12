#!/usr/bin/env python3
# SPDX-License-Identifier: MIT
# Copyright (c) 2026 Kevin Wuytack (Holdvero)
"""Check an Open Evidence Pack (Excel, YAML or JSON) against the format.

Usage:
    python tools/validate.py FILE [FILE ...] [--strict] [--today YYYY-MM-DD] [--to-json OUT.json]

Works with the English and the Dutch Excel workbook, and with YAML or JSON files.
Extra columns you add to the workbook are kept as x_ fields and do not cause errors.

Errors (exit code 1):
    - the pack does not match format/evidence-pack.schema.json
    - two answers, evidence entries or open items share the same id
    - an answer refers to evidence or an open item that does not exist
    - an open item refers to a question that does not exist
    - a next review date lies before the last review date
    - a hash does not have the length its algorithm requires

Warnings (exit code 0, or 1 with --strict):
    - a review date, evidence review date or open-item target date has passed
    - an evidenced answer relies on evidence that is not current
    - an answer is still partial or open while all its open items are closed

--to-json writes the pack as JSON, for example to convert an Excel workbook.
"""
import argparse
import datetime as dt
import json
import re
import sys
from pathlib import Path

try:
    import yaml
    from jsonschema import Draft202012Validator, FormatChecker
except ImportError:
    sys.exit("Missing packages. Install them with: pip install jsonschema pyyaml openpyxl")

SCHEMA_PATH = Path(__file__).resolve().parent.parent / "format" / "evidence-pack.schema.json"
FORMAT_VERSION = "0.2"

# ---------------------------------------------------------------- Excel mapping

SHEETS = {
    "answers": ("Answers", "Antwoorden"),
    "evidence": ("Evidence register", "Bewijsregister"),
    "open_items": ("Open items", "Open punten"),
}

# Dutch column headers. English headers map automatically ("Question ID" -> question_id).
HEADERS_NL = {
    "vraag-id": "question_id", "categorie": "category", "vraag": "question",
    "kort antwoord": "short_answer", "goedgekeurd antwoord": "approved_answer",
    "status antwoord": "answer_status", "eigenaar antwoord": "answer_owner",
    "laatst herzien": "last_reviewed", "volgende herziening": "next_review",
    "bewijs": "evidence", "open punten": "open_item_ids", "kaderverwijzingen": "framework_refs",
    "notities": "notes",
    "bewijs-id": "evidence_id", "titel": "title", "soort bewijs": "evidence_type",
    "versie": "version", "versiedatum": "version_date", "eigenaar": "owner",
    "herzien voor": "review_due", "vindplaats": "locator", "deelbaar": "shareable",
    "status": "status", "opslaglocatie": "storage_reference", "hash-algoritme": "hash_algorithm",
    "hash": "hash",
    "open-punt-id": "open_item_id", "omschrijving": "description",
    "betreft vragen": "related_question_ids", "streefdatum": "target_date",
    "afhandeling": "resolution_note",
}

# Dutch dropdown values per field. English values are the canonical values.
VALUES_NL = {
    "short_answer": {"ja": "yes", "nee": "no", "deels": "partly", "beschreven": "described", "n.v.t.": "not_applicable"},
    "answer_status": {"onderbouwd": "evidenced", "deels onderbouwd": "partial", "open punt": "open_item",
                      "nee met uitleg": "explained_no", "n.v.t.": "not_applicable"},
    "evidence_type": {"beleid": "policy", "procedure": "procedure", "registratie": "record",
                      "documentatie": "documentation", "configuratie": "configuration",
                      "systeemexport": "system_generated", "schermafdruk": "screenshot",
                      "testresultaat": "test_result", "auditrapport": "audit_report",
                      "certificaat": "certificate", "contract": "contract", "overig": "other"},
    "shareable": {"openbaar": "public", "op aanvraag": "on_request", "onder nda": "under_nda",
                  "alleen intern": "internal_only"},
    "evidence_status": {"actueel": "current", "concept": "draft", "vervangen": "superseded",
                        "ingetrokken": "withdrawn"},
    "open_item_status": {"open": "open", "bezig": "in_progress", "afgerond": "closed"},
}

LIST_FIELDS = {"open_item_ids", "framework_refs", "related_question_ids"}
KNOWN = {
    "answers": {"question_id", "category", "question", "short_answer", "approved_answer", "answer_status", "answer_owner",
                "last_reviewed", "next_review", "evidence", "open_item_ids", "framework_refs", "notes", "language"},
    "evidence": {"evidence_id", "title", "evidence_type", "version", "version_date", "owner", "review_due", "locator",
                 "shareable", "status", "storage_reference", "hash_algorithm", "hash", "notes"},
    "open_items": {"open_item_id", "description", "related_question_ids", "owner", "target_date", "status", "resolution_note"},
}
DATE_FIELDS = {"last_reviewed", "next_review", "version_date", "review_due", "target_date"}


def header_to_field(text):
    key = str(text).strip().lower()
    if key in HEADERS_NL:
        return HEADERS_NL[key]
    return re.sub(r"[^a-z0-9]+", "_", key).strip("_")


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, (dt.datetime, dt.date)):
        return value.isoformat()[:10]
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def split_list(text, commas=True):
    parts = re.split(r"[;\n]" + ("|," if commas else ""), text)
    return [p.strip() for p in parts if p.strip()]


def map_value(section, field, value):
    key = {"evidence": "evidence_status", "open_items": "open_item_status"}.get(section, field) \
        if field == "status" else field
    table = VALUES_NL.get(key, {})
    return table.get(value.lower(), value)


def read_sheet(ws, section):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    fields = []
    for h in rows[0]:
        field = header_to_field(h) if h is not None and str(h).strip() else None
        if field and field not in KNOWN[section]:
            field = "x_" + field  # your own extra column: kept as an extension field
        fields.append(field)
    records = []
    for row in rows[1:]:
        if all(cell_text(v) == "" for v in row):
            continue
        rec = {}
        for field, value in zip(fields, row):
            text = cell_text(value)
            if not field or text == "":
                continue
            if field == "evidence":
                refs = []
                for item in split_list(text, commas=False):
                    ev_id, _, locator = item.partition(" ")
                    ref = {"evidence_id": ev_id.strip()}
                    if locator.strip():
                        ref["locator"] = locator.strip()
                    refs.append(ref)
                rec[field] = refs
            elif field in LIST_FIELDS:
                rec[field] = split_list(text, commas=field != "framework_refs")
            else:
                rec[field] = map_value(section, field, text)
        records.append(rec)
    return records


def load_xlsx(path):
    try:
        from openpyxl import load_workbook
    except ImportError:
        sys.exit("Reading Excel needs openpyxl. Install it with: pip install openpyxl")
    wb = load_workbook(path, data_only=True)

    def named(name):
        dn = wb.defined_names.get(name)
        if dn is None:
            return ""
        for sheet, coord in dn.destinations:
            return cell_text(wb[sheet][coord.replace("$", "")].value)
        return ""

    data = {"format_version": FORMAT_VERSION, "organisation": {"name": named("org_name")}}
    if named("org_website"):
        data["organisation"]["website"] = named("org_website")
    if named("as_of"):
        data["as_of"] = named("as_of")
    data["default_language"] = "nl" if "Antwoorden" in wb.sheetnames else "en"
    if not data["organisation"]["name"]:
        del data["organisation"]["name"]
    for section, names in SHEETS.items():
        for name in names:
            if name in wb.sheetnames:
                data[section] = read_sheet(wb[name], section)
                break
    return data


# ---------------------------------------------------------------- YAML / JSON

def normalise(value):
    """YAML turns unquoted dates into date objects; the schema expects strings."""
    if isinstance(value, dict):
        return {str(k): normalise(v) for k, v in value.items()}
    if isinstance(value, list):
        return [normalise(v) for v in value]
    if isinstance(value, (dt.date, dt.datetime)):
        return value.isoformat()[:10]
    return value


def load(path):
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        return load_xlsx(path)
    text = path.read_text(encoding="utf-8")
    if suffix == ".json":
        return json.loads(text)
    return normalise(yaml.safe_load(text))


# ---------------------------------------------------------------- checks

def as_date(value):
    try:
        return dt.date.fromisoformat(value)
    except (TypeError, ValueError):
        return None


def location(error_path):
    out = ""
    for p in error_path:
        out += f"[{p}]" if isinstance(p, int) else (f".{p}" if out else p)
    return out or "(pack)"


def duplicates(values):
    return sorted({v for v in values if v is not None and values.count(v) > 1})


def check(data, today):
    errors, warnings = [], []
    schema = json.loads(SCHEMA_PATH.read_text(encoding="utf-8"))
    validator = Draft202012Validator(schema, format_checker=FormatChecker())
    for err in sorted(validator.iter_errors(data), key=lambda e: [str(p) for p in e.absolute_path]):
        errors.append(f"{location(err.absolute_path)}: {err.message}")
    if not isinstance(data, dict):
        return errors, warnings

    def items(key):
        return [x for x in (data.get(key) or []) if isinstance(x, dict)]

    answers, evidence, open_items = items("answers"), items("evidence"), items("open_items")
    q_ids = [a.get("question_id") for a in answers]
    ev_ids = [e.get("evidence_id") for e in evidence]
    oi_ids = [o.get("open_item_id") for o in open_items]
    for label, ids in (("answers", q_ids), ("evidence", ev_ids), ("open_items", oi_ids)):
        for dup in duplicates(ids):
            errors.append(f"{label}: id '{dup}' is used more than once")

    ev_by_id = {e.get("evidence_id"): e for e in evidence}
    oi_by_id = {o.get("open_item_id"): o for o in open_items}
    hash_len = {"sha256": 64, "sha384": 96, "sha512": 128}

    for e in evidence:
        label = f"evidence {e.get('evidence_id', '?')}"
        algo, digest = e.get("hash_algorithm"), e.get("hash")
        if algo in hash_len and isinstance(digest, str) and len(digest) != hash_len[algo]:
            errors.append(f"{label}: a {algo} hash has {hash_len[algo]} characters, this one has {len(digest)}")
        due = as_date(e.get("review_due"))
        if due and due < today and e.get("status") == "current":
            warnings.append(f"{label}: review was due on {due}")

    for a in answers:
        label = f"answer {a.get('question_id', '?')}"
        for ref in a.get("evidence") or []:
            ev_id = ref.get("evidence_id") if isinstance(ref, dict) else None
            if ev_id not in ev_by_id:
                errors.append(f"{label}: evidence '{ev_id}' is not in the evidence register")
            elif a.get("answer_status") == "evidenced" and ev_by_id[ev_id].get("status") != "current":
                warnings.append(f"{label}: relies on evidence '{ev_id}' with status '{ev_by_id[ev_id].get('status')}'")
        for oi in a.get("open_item_ids") or []:
            if oi not in oi_by_id:
                errors.append(f"{label}: open item '{oi}' does not exist")
        linked = [oi_by_id[o] for o in (a.get("open_item_ids") or []) if o in oi_by_id]
        if a.get("answer_status") in ("partial", "open_item") and linked and all(o.get("status") == "closed" for o in linked):
            warnings.append(f"{label}: all its open items are closed; update the answer and its status")
        last, nxt = as_date(a.get("last_reviewed")), as_date(a.get("next_review"))
        if last and nxt and nxt < last:
            errors.append(f"{label}: next_review {nxt} lies before last_reviewed {last}")
        if nxt and nxt < today:
            warnings.append(f"{label}: review was due on {nxt}")

    for o in open_items:
        label = f"open item {o.get('open_item_id', '?')}"
        for q in o.get("related_question_ids") or []:
            if q not in q_ids:
                errors.append(f"{label}: question '{q}' does not exist")
        target = as_date(o.get("target_date"))
        if target and target < today and o.get("status") != "closed":
            warnings.append(f"{label}: target date {target} has passed")

    return errors, warnings


def main():
    parser = argparse.ArgumentParser(description="Check an Open Evidence Pack against the format.")
    parser.add_argument("files", nargs="+", type=Path)
    parser.add_argument("--strict", action="store_true", help="treat warnings as errors")
    parser.add_argument("--today", type=dt.date.fromisoformat, default=dt.date.today(),
                        help="date used to check review and target dates (default: today)")
    parser.add_argument("--to-json", type=Path, metavar="OUT", help="write the (single) pack as JSON")
    args = parser.parse_args()
    if args.to_json and len(args.files) != 1:
        parser.error("--to-json works with one input file")

    failed = False
    for path in args.files:
        try:
            data = load(path)
        except Exception as exc:  # unreadable file, bad YAML, impossible date, ...
            print(f"FAIL  {path}\n  cannot read file: {exc}")
            failed = True
            continue
        errors, warnings = check(data, args.today)
        bad = bool(errors) or (args.strict and bool(warnings))
        failed = failed or bad
        counts = {k: len(data.get(k) or []) for k in ("answers", "evidence", "open_items")} if isinstance(data, dict) else {}
        summary = ", ".join(f"{v} {k.replace('_', ' ')}" for k, v in counts.items())
        print(f"{'FAIL' if bad else 'OK  '}  {path}  ({summary}; {len(errors)} errors, {len(warnings)} warnings)")
        for e in errors:
            print(f"  error:   {e}")
        for w in warnings:
            print(f"  warning: {w}")
        if args.to_json:
            args.to_json.write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
            print(f"  written: {args.to_json}")

    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
